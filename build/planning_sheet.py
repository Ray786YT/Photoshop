"""Planning sheet in plain, first-person language, matching the student's steps."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

NAME = 'Rayyan Ali'
ACCENT = RGBColor(0x2E, 0x5B, 0x2E)

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Inches(0.8)
    s.left_margin = s.right_margin = Inches(0.9)
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(8)

def head(t, size=16):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t); r.bold = True; r.font.size = Pt(size)
def q(t):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(14)
    r = p.add_run(t); r.bold = True; r.font.color.rgb = ACCENT
def a(t, bullet=False):
    p = doc.add_paragraph(style='List Bullet' if bullet else None); p.add_run(t)
    if not bullet: p.paragraph_format.left_indent = Inches(0.25)

head('MAGAZINE COVER PLANNING SHEET')
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run(f'Student Name: {NAME}').bold = True

q('1. What type of magazine cover are you creating?')
a('Technology / AI. My magazine is about how technology is changing the way people '
  'live and talk to each other.')

q('2. What is the name of your magazine?')
a('PROMPT. A prompt is what you type into an AI, but it also means helping someone '
  'start talking. I liked that it has both meanings.')

q('3. In one sentence, what is the main message of your magazine cover (what is the '
  'main cover line)?')
a('"FINISH MY SENTENCE." - we are letting AI finish our sentences for us, so we are '
  'losing our own voice.')

q('4a. What is your central image going to look like?')
a('A black and white photo of a woman looking straight at the camera. Her mouth is '
  'removed so there is just skin there. Where her mouth used to be, I will put the '
  'word suggestion bar from a phone keyboard. It will say "I\'m | fine | thanks", '
  'with "fine" highlighted in green like the phone already picked it for her.')

q('4b. How are you going to manipulate the image? Be detailed.')
a('Remove the mouth: select it with the Lasso tool, use Content-Aware Fill, then clean '
  'it up with the Spot Healing Brush and Clone Stamp. Content-Aware Fill is my new skill - '
  'I will learn it from a YouTube tutorial.', True)
a('Make the grey background almost black with a Curves adjustment layer.', True)
a('Darken the edges with a black Vignette layer (a big soft Eraser in the middle).', True)
a('Turn the photo dark blue with a Gradient Map adjustment layer.', True)
a('Use the Gradient tool to fade the bottom to black so the bottom text is easy to read.', True)
a('Make the suggestion bar with the Rounded Rectangle tool and the Type tool, and give '
  'it a drop shadow.', True)
a('Use downloaded fonts: Manrope and Roboto from Google Fonts.', True)

q('4c. How is this manipulation enhancing or helping you convey the message?')
a('Taking away her mouth shows she can\'t speak for herself anymore. Putting the '
  'keyboard bar where her mouth should be shows the phone is talking for her. I chose '
  '"I\'m fine thanks" because people say it automatically even when it isn\'t true. '
  'The dark blue makes it feel cold and robotic. The green is the only bright colour, '
  'so your eye goes to the bar first and then down to "SENTENCE."')

q('5. Did Ms. Bensusan discuss your idea with you and give you feedback?')
a('[ Write what Ms. Bensusan said here, and anything you changed because of it. ]')

doc.add_page_break()
head('ELEMENTS AND PRINCIPLES OF DESIGN', 14)

para_note = doc.add_paragraph()
para_note.add_run('For each one: where it is on my cover, how I used it, and why.').italic = True

def ep(title, text):
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(title + ' - ').bold = True
    p.add_run(text)

q('Elements of Design')
ep('Line', 'I put a thin green line right under the PROMPT title to separate the title from '
   'the photo. There are also two thin grey lines inside the keyboard bar that split it into '
   'three boxes, the same as a real phone keyboard, so people recognise it straight away.')
ep('Shape', 'The suggestion bar is a long rounded rectangle, the same shape as the one on a '
   'phone, so it looks like real technology. It goes straight across her face where her mouth '
   'should be, so its flat, straight shape cuts across the round shape of her face and stands out.')
ep('Value', 'I used a Curves adjustment layer to make the background almost black while keeping '
   'her face bright. The big difference between light and dark makes her face the brightest '
   'thing on the page, so it is the first thing you look at. The dark background also makes the '
   'white text easy to read. I also darkened the edges with a Vignette layer so the light '
   'stays in the middle.')
ep('Colour', 'Most of the cover is dark blue from the Gradient Map, which feels cold and '
   'robotic, like technology. The only bright colour is green, and I only used it on the '
   'important parts: the selected word "fine", the small headings, the line under the title, '
   'and "SENTENCE." Because it is the only bright colour, the green shows you where to look.')
ep('Texture', 'I kept the real skin texture on her face, like the pores and small lines, so she '
   'looks like a real person and not a drawing. When I removed her mouth, I cleaned it up with '
   'the Spot Healing Brush so the skin texture there matches the rest of her face. If that spot '
   'was smooth and blurry, the edit would look fake.')
ep('Space', 'I left empty dark space on both sides of her head and put my three smaller cover '
   'lines there, so the text never covers her face. The empty space also makes the cover feel '
   'quiet and a bit lonely, which fits the idea of someone who has lost their voice.')
ep('Form', 'The shadows on her nose, cheeks and jaw make her face look 3D and real. The keyboard '
   'bar is completely flat, so it looks like something stuck onto her, which shows that the '
   'technology does not belong there.')

q('Principles of Design')
ep('Balance', 'Her face is right in the centre and the title is centred at the top, so the cover '
   'is balanced from left to right. I balanced the text around her too: two cover lines on the '
   'left and one on the right, and the big main cover line at the bottom left with the barcode '
   'at the bottom right.')
ep('Contrast', 'I used contrast in three ways: a bright face against a dark background, bright '
   'green against dark blue, and a huge heavy title against the small thin selling line above '
   'it. The strongest contrast is the light bar and green box on the darker part of her face, '
   'so that is what gets noticed first.')
ep('Emphasis', 'The keyboard bar is the focal point. It is in the middle of the page, it has the '
   'only bright green box, and it is where her mouth should be, which looks strange and makes '
   'you look twice. That is on purpose, because the bar is my main message.')
ep('Movement', 'Your eye starts at PROMPT at the top, moves down her face to her eyes, then to '
   'the bar, then down to "FINISH MY SENTENCE." at the bottom. The green on the title line, the '
   'bar and "SENTENCE." works like stepping stones that lead your eye down the page.')
ep('Repetition', 'I repeated the same green four times. I also set up all three small cover '
   'lines the exact same way: a small green word on top and bold white text underneath. '
   'Repeating things like a real magazine does makes the cover look planned and professional.')
ep('Unity', 'I only used two fonts, Manrope and Roboto, and the same few colours everywhere. '
   'Turning the photo dark blue made it match the dark background, so the photo and the text '
   'look like one design instead of separate pieces put together.')
ep('Rhythm', 'The three small cover lines follow the same pattern (green heading, then three '
   'short lines of white text) and are spaced out down the sides. Seeing the same pattern '
   'again and again gives the cover a steady beat as you read it.')

q('Checklist')
for t in ['8.5" x 11" at 300 ppi, in colour.',
          'Downloaded fonts: Manrope and Roboto.',
          'Masthead, selling line, main cover line, three more cover lines, dateline with my '
          'name, and a barcode.',
          'Layers organised into named folders.',
          'Selection tool: Lasso. Editing tools: Clone Stamp, Spot Healing Brush, resizing.',
          'Colour tools: Curves and Gradient Map adjustment layers.',
          'New skill learned on my own: Content-Aware Fill.']:
    a(t, True)

doc.save('/home/user/Photoshop/out/PROMPT_planning_sheet.docx')
print('planning sheet saved')
