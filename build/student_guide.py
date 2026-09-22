"""Finish-It-Yourself guide for PROMPT_STARTER.psd, written at Grade 11 level."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor

GREEN = RGBColor(0x2E, 0x6B, 0x1E)
GREY = RGBColor(0x55, 0x5B, 0x55)

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Inches(0.75)
    s.left_margin = s.right_margin = Inches(0.9)
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11)

def h1(t):
    r = doc.add_paragraph().add_run(t); r.bold = True; r.font.size = Pt(18)
def h2(t):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(14)
    r = p.add_run(t); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = GREEN
def para(t, italic=False, color=None):
    r = doc.add_paragraph().add_run(t); r.italic = italic
    if color: r.font.color.rgb = color
def step(t):
    doc.add_paragraph(t, style='List Number')
def bullet(t):
    doc.add_paragraph(t, style='List Bullet')
def shot(t):
    p = doc.add_paragraph(); r = p.add_run('SCREENSHOT: ' + t); r.bold = True
def why(t):
    para('Why: ' + t, italic=True, color=GREY)

h1('PROMPT Cover - Finish It Yourself')
para('Open PROMPT_STARTER.psd in Photoshop. The text and layout are already done. '
     'You do the photo editing - the part your teacher is marking. About 1 class '
     'period. Take a screenshot at every camera step for your process work.')
para('Can\'t see the layers? Window > Layers.', italic=True)
shot('the starter file open, with the Layers panel showing.')

h2('Step 1 - Remove the mouth (NEW SKILL: Content-Aware Fill)')
para('This is your "independently learned" skill. Before you start, watch a short '
     'YouTube video: search "Photoshop content aware fill tutorial".')
step('Click the eye next to the 03 AUTOCOMPLETE BAR folder to hide it, so you can see the mouth.')
step('Click the layer "Portrait - original photo". Press Ctrl+J (Mac: Cmd+J) to copy it. '
     'Double-click the copy\'s name and call it "Portrait - no mouth". The one underneath '
     'is your BEFORE - never edit it.')
step('Pick the Lasso tool (L). Draw a loop around the mouth, a little bigger than the lips.')
step('Go to Edit > Content-Aware Fill, then click OK. (Or Edit > Fill, Contents: Content-Aware.)')
step('Press Ctrl+D (Cmd+D) to deselect. Zoom in.')
step('Fix any smudges with the Spot Healing Brush (J) - just paint over them. For bigger messy '
     'spots use the Clone Stamp (S): hold Alt (Mac: Option) and click clean cheek skin, then '
     'paint over the mess.')
shot('before and after - click the eye on "Portrait - no mouth" off and on.')
why('taking away her mouth shows she has lost her own voice.')

h2('Step 2 - Make the background dark (Curves)')
step('Click "Portrait - no mouth" so it is selected.')
step('Layer > New Adjustment Layer > Curves > OK.')
step('In the Properties panel, click the middle of the diagonal line and drag it DOWN. '
     'Then click near the top-right of the line and drag it up a little.')
step('Keep adjusting until the grey background is almost black but her face is still bright.')
para('If you want exact numbers: one point at Input 145 / Output 58, another at '
     'Input 205 / Output 196. Close enough is fine.', italic=True, color=GREY)
shot('the Curves panel open.')
why('a dark background makes the face and the white text stand out (contrast).')

h2('Step 3 - Add colour (Gradient Map)')
step('Layer > New Adjustment Layer > Gradient Map > OK.')
step('In Properties, click the gradient bar to open the Gradient Editor.')
step('Click the LEFT colour stop under the bar, set it to #06080E (very dark blue).')
step('Click the RIGHT colour stop, set it to #F4F6F2 (off-white).')
step('Click just under the bar to add three more stops. For each one, type its Location '
     'and pick its colour: 30% #1A2230 (dark navy), 58% #6C767E (blue-grey), '
     '82% #CDD2CE (light grey). Click OK.')
shot('the Gradient Editor open.')
why('the cold blue feels robotic and techy, and it makes the green bar the only bright '
    'colour, so your eye goes straight to it.')

h2('Step 4 - Save both files')
step('Check Image > Image Size says 8.5 x 11 inches at 300 Pixels/Inch.')
step('File > Save. This keeps all your layers (.psd).')
step('File > Save a Copy (or Export > Export As) > JPEG, highest quality.')
shot('your finished Layers panel.')

h2('If your teacher asks how you did it')
bullet('Mouth: "I lassoed it and used Content-Aware Fill, then cleaned it up with the '
       'Spot Healing Brush and Clone Stamp."')
bullet('Dark background: "A Curves adjustment layer."')
bullet('Blue colour: "A Gradient Map adjustment layer."')
bullet('New skill: "I learned Content-Aware Fill from a YouTube tutorial. You select '
       'something and Photoshop fills it in using the pixels around it."')
bullet('Fonts: "Manrope and Roboto, downloaded from Google Fonts."')

h2('Optional extras (only if you want)')
bullet('Filter > Noise > Add Noise (about 2%, Monochromatic) on the healed area, so the '
       'skin texture matches. This is a professional retoucher trick, not a normal '
       'Grade 11 step - skip it if you like.')
bullet('Want to change any words? Hide that text layer, use the Type tool (T) and '
       'type your own. Install Manrope and Roboto from fonts.google.com first.')

doc.save('/home/user/Photoshop/out/PROMPT_FINISH_IT_YOURSELF_guide.docx')
print('guide saved')
